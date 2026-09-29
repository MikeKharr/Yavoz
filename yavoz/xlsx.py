"""Чтение и заполнение книги: столбцы B, C, D, S, T — вход, U и V — выход."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.worksheet import Worksheet

from .rank import Персона, нормализовать

ФИО = ("B", "C", "D")
МЕСТО, ДОЛЖНОСТЬ = "S", "T"
ПО_ЗАПРОСУ, В_ИНТЕРНЕТЕ = "U", "V"


class КнигаНеПодходит(RuntimeError):
    pass


@dataclass
class Строка:
    номер: int
    персона: Персона
    u_пуст: bool
    v_пуст: bool


def формула_запроса(номер: int) -> str:
    """Та же формула, что в образце, с подставленным номером строки."""
    n = номер
    return ('=HYPERLINK("https://yandex.ru/search/?text=" & '
            f'B{n} & " " & C{n} & " " & D{n} & " " & S{n} & " " & T{n})')


def найти_лист(книга) -> Worksheet:
    """Лист данных — тот, где в шапке столбца B стоит «Фамилия».

    Опираемся только на столбцы из формулы: остальные могут быть пустыми.
    """
    подходящие = []
    for имя in книга.sheetnames:
        ws = книга[имя]
        шапка = нормализовать(str(ws["B1"].value or ""))
        if "фамилия" in шапка:
            подходящие.append(ws)
    if not подходящие:
        raise КнигаНеПодходит("не нашёл лист, где в B1 стоит «Фамилия»")
    # если таких несколько — берём самый длинный, это база, а не справочник
    return max(подходящие, key=lambda ws: ws.max_row)


def строки(ws: Worksheet, предел: int = 0) -> list[Строка]:
    итог: list[Строка] = []
    for n in range(2, ws.max_row + 1):
        фамилия = ws[f"B{n}"].value
        if фамилия is None or not str(фамилия).strip():
            continue
        персона = Персона(
            фамилия=str(фамилия).strip(),
            имя=str(ws[f"C{n}"].value or "").strip(),
            отчество=str(ws[f"D{n}"].value or "").strip(),
            организация=str(ws[f"{МЕСТО}{n}"].value or "").strip(),
            должность=str(ws[f"{ДОЛЖНОСТЬ}{n}"].value or "").strip(),
        )
        итог.append(Строка(
            номер=n,
            персона=персона,
            u_пуст=not str(ws[f"{ПО_ЗАПРОСУ}{n}"].value or "").strip(),
            v_пуст=not str(ws[f"{В_ИНТЕРНЕТЕ}{n}"].value or "").strip(),
        ))
        if предел and len(итог) >= предел:
            break
    return итог


def записать_u(ws: Worksheet, номер: int) -> None:
    ws[f"{ПО_ЗАПРОСУ}{номер}"] = формула_запроса(номер)


def записать_v(ws: Worksheet, номер: int, url: str, пометка: str) -> None:
    ячейка = ws[f"{В_ИНТЕРНЕТЕ}{номер}"]
    ячейка.value = url
    # Уверенность живёт примечанием: в значении ячейки она сломала бы ссылку,
    # а соседние столбцы заняты данными владельца.
    ячейка.comment = Comment(пометка, "Yavoz")


def открыть(путь: Path):
    return load_workbook(путь)


def сохранить(книга, путь: Path) -> Path:
    путь.parent.mkdir(parents=True, exist_ok=True)
    книга.save(путь)
    return путь
